"""US-02 et US-15 — mentions obligatoires BTP et régime de TVA.

Le test central est `test_en_franchise_aucune_tva_hors_mention_legale` : il lit
le texte réellement produit dans le .docx, retire la mention légale, et exige
qu'il ne reste **aucune** occurrence de « TVA ». C'est ce que le critère de
validation d'US-02 demande, mot pour mot.
"""
from __future__ import annotations

import datetime as dt
import io

import pytest
from docx import Document

from app.devis import Devis, ajouter_ligne
from app.export import MentionsIncompletes, generer_docx
from app.profil import MENTION_293B, Client, Decennale, ProfilArtisan, RegimeTVA


@pytest.fixture
def profil() -> ProfilArtisan:
    return ProfilArtisan(
        raison_sociale="Karim Rénovation SARL", siret="81234567800017",
        adresse="12 rue des Lilas, 69003 Lyon", rcs="RCS Lyon 812 345 678",
        ape_naf="4334Z", tva_intracommunautaire="FR40812345678",
        decennale=Decennale(assureur="AXA", numero_contrat="DEC-99812",
                            couverture_geographique="France métropolitaine",
                            valide_jusqu_au=dt.date(2027, 6, 30)),
    )


@pytest.fixture
def client() -> Client:
    return Client(nom="Mme Dupont", adresse_chantier="4 impasse du Parc, 69100 Villeurbanne")


@pytest.fixture
def devis() -> Devis:
    d = Devis()
    ajouter_ligne(d, "PEINT_FINITION", "surface", 22.0)
    return d


def texte_du_docx(contenu: bytes) -> str:
    """Tout le texte du document : paragraphes et cellules de tableau."""
    doc = Document(io.BytesIO(contenu))
    morceaux = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for ligne in table.rows:
            morceaux.extend(cellule.text for cellule in ligne.cells)
    return "\n".join(morceaux)


# --- Cas dégradé : une donnée légale manquante bloque l'export --------------

def test_export_refuse_si_une_mention_legale_manque(devis, client):
    incomplet = ProfilArtisan(raison_sociale="Karim Rénovation SARL")
    with pytest.raises(MentionsIncompletes) as erreur:
        generer_docx(devis, [], incomplet, client, "3 semaines")
    assert "votre numéro SIRET" in erreur.value.manquants


def test_le_delai_d_execution_est_une_mention_obligatoire(devis, profil, client):
    with pytest.raises(MentionsIncompletes) as erreur:
        generer_docx(devis, [], profil, client, "")
    assert erreur.value.manquants == ["le délai d'exécution des travaux"]


# --- Nominal : régime réel --------------------------------------------------

def test_le_document_porte_toutes_les_mentions_obligatoires(devis, profil, client):
    texte = texte_du_docx(
        generer_docx(devis, ["salon"], profil, client, "3 semaines")
    )
    for attendu in ("81234567800017", "RCS Lyon 812 345 678", "4334Z",
                    "FR40812345678", "AXA", "DEC-99812", "Mme Dupont",
                    "4 impasse du Parc", "3 semaines", "Bon pour accord"):
        assert attendu in texte, f"mention absente du devis : {attendu}"


def test_en_regime_reel_la_ventilation_tva_est_presente(devis, profil, client):
    texte = texte_du_docx(generer_docx(devis, [], profil, client, "3 semaines"))
    assert "Total TTC" in texte
    assert "Total TVA" in texte
    assert MENTION_293B not in texte


# --- Franchise en base, art. 293 B du CGI -----------------------------------

def test_en_franchise_aucune_tva_hors_mention_legale(devis, profil, client):
    profil.regime_tva = RegimeTVA.FRANCHISE_293B
    texte = texte_du_docx(generer_docx(devis, [], profil, client, "3 semaines"))

    assert MENTION_293B in texte
    reste = texte.replace(MENTION_293B, "")
    assert "TVA" not in reste, (
        "le devis d'un artisan en franchise ne doit porter aucune TVA "
        f"hors la mention légale ; trouvé dans : {reste!r}"
    )


def test_en_franchise_ni_total_ttc_ni_colonne_tva(devis, profil, client):
    profil.regime_tva = RegimeTVA.FRANCHISE_293B
    texte = texte_du_docx(generer_docx(devis, [], profil, client, "3 semaines"))
    assert "TTC" not in texte
    assert "Total HT" not in texte
    assert "Total :" in texte


def test_en_franchise_le_numero_intracom_n_est_pas_exige(devis, profil, client):
    """Un artisan en franchise n'en a pas nécessairement — l'exiger bloquerait
    un export pourtant conforme."""
    profil.regime_tva = RegimeTVA.FRANCHISE_293B
    profil.tva_intracommunautaire = ""
    generer_docx(devis, [], profil, client, "3 semaines")  # ne lève pas


def test_l_acompte_en_franchise_porte_sur_le_montant_net(devis, profil, client):
    profil.regime_tva = RegimeTVA.FRANCHISE_293B
    texte = texte_du_docx(generer_docx(devis, [], profil, client, "3 semaines"))
    attendu = round(devis.total_ht * 0.30, 2)
    assert f"Acompte demandé (30 %) : {attendu:.2f} €" in texte


# --- Décennale --------------------------------------------------------------

def test_decennale_perimee_est_signalee_sans_bloquer(devis, profil, client):
    profil.decennale.valide_jusqu_au = dt.date(2024, 1, 31)
    texte = texte_du_docx(
        generer_docx(devis, [], profil, client, "3 semaines",
                     aujourd_hui=dt.date(2026, 9, 8))
    )
    assert "expirée depuis le 31/01/2024" in texte
    assert "Bon pour accord" in texte  # l'émission n'est pas bloquée


def test_echeance_inconnue_n_est_pas_traitee_comme_perimee():
    assert Decennale().perimee() is False


# --- Structure, pour l'export PDF balisé ------------------------------------

def test_le_document_a_des_titres_et_un_tableau_a_en_tete(devis, profil, client):
    doc = Document(io.BytesIO(
        generer_docx(devis, [], profil, client, "3 semaines")))
    niveaux = [p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Heading 1" in niveaux and "Heading 2" in niveaux
    assert doc.tables[0].rows[0].cells[0].text == "Poste"
