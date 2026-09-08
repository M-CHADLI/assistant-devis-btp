"""US-03 — reprendre la mise en page d'un ancien devis.

Ces tests portent les critères de la story, dans l'ordre où ils comptent :

1. **Rien du document déposé n'est conservé** — c'est la promesse la plus forte
   du module, et celle qu'un relecteur doit pouvoir vérifier en un test.
2. Un dépôt qu'on sait refuser ne sort jamais de la machine.
3. Un échec de lecture, quelle qu'en soit la cause, ne dégrade pas l'export :
   la mise en page par défaut reste disponible et le motif est nommé.
4. Un document qui n'est pas un devis est refusé avec un message qui dit ce
   qui a été cherché sans être trouvé.
5. Le gabarit réordonne et renomme, il ne retire jamais.
6. Le fichier n'est jamais soumis au modèle.

Aucun de ces tests n'a besoin de réseau ni de clé : le transport est injecté.
Ce que ces tests **ne prouvent pas** est écrit à la fin du fichier.
"""
from __future__ import annotations

import ast
import io
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from tests.conftest import authentifier

from app import gabarit as mod
from app import main, pipeline
from app.gabarit import (BLOCS, DepotInvalide, Gabarit, Indisponible,
                         PasUnDevis, analyser, extraire, lire_markdown,
                         valider_depot, vocabulaire)
from app.persistance import Registre

RACINE = Path(__file__).resolve().parents[1]

PDF = b"%PDF-1.7\n" + b"0" * 400
JPEG = b"\xff\xd8\xff\xe0" + b"0" * 400
PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 400

# Un devis d'artisan tel que l'OCR le rendrait : titres en gras, intitulés
# maison, et — c'est le point — des données personnelles d'un *ancien* client.
DEVIS_MARKDOWN = """
# DEVIS N° 2024-118

**Émetteur**
SARL Peinture Duverger — 4 place du Marché, 33000 Bordeaux
SIRET 512 345 678 00019

**Destinataire**
Mme Solange Petitjean
17 impasse des Glycines, 33270 Floirac

**Prestations**

| Désignation | Qté | PU HT | Total HT |
|---|---|---|---|
| Peinture murs salon | 42 m² | 18,00 € | 756,00 € |
| Ragréage sol | 21 m² | 24,50 € | 514,50 € |

**Totaux**
Total HT : 1 270,50 €
TVA 10 % : 127,05 €
Total TTC : 1 397,55 €

**Conditions**
Acompte de 30 % à la commande. Assurance décennale AXA n° 4471822.

**Bon pour accord**
Date et signature du client :
"""


def reponse_ocr(markdown: str, confiance: float | None = 0.93) -> dict:
    page: dict = {"index": 0, "markdown": markdown}
    if confiance is not None:
        page["confidence_scores"] = {"average_page_confidence_score": confiance}
    return {"model": mod.MODELE_OCR, "pages": [page],
            "usage_info": {"pages_processed": 1}}


# --------------------------------------------------------------------------
# 1. Rien du document déposé n'est conservé
# --------------------------------------------------------------------------

def test_le_gabarit_ne_contient_que_le_vocabulaire_du_module():
    """La promesse, vérifiée mot à mot.

    Si un jour l'analyse se met à retenir du texte du document, ce test tombe
    — c'est exactement ce qu'on lui demande.
    """
    gab = analyser(DEVIS_MARKDOWN)
    connus = vocabulaire()
    for bloc in gab.blocs:
        assert bloc in connus
    for bloc, synonyme in gab.intitules.items():
        assert bloc in BLOCS
        assert synonyme in BLOCS[bloc]


def test_aucune_donnee_personnelle_de_l_ancien_client_ne_survit():
    gab = analyser(DEVIS_MARKDOWN)
    serialise = str(gab.en_dict())
    for fuite in ("Petitjean", "Solange", "Glycines", "Floirac", "Duverger",
                  "512 345 678", "1 397,55", "4471822", "2024-118", "AXA"):
        assert fuite not in serialise


def test_aucun_chiffre_ne_peut_entrer_dans_un_gabarit():
    """Montants, dates, numéros de contrat : aucun ne franchit l'analyse."""
    gab = analyser(DEVIS_MARKDOWN)
    assert not any(c.isdigit() for c in str(gab.en_dict()))


def test_un_intitule_qui_porterait_un_nom_ne_le_transmet_pas():
    """« Devis pour Mme Petitjean » reconnaît le bloc sans en garder le texte."""
    gab = analyser("## Devis n° 44 pour Mme Petitjean\n"
                   "## Prestations\n## Totaux\n")
    assert "identification" in gab.blocs
    assert "Petitjean" not in str(gab.en_dict())
    # Le synonyme reconnu n'étant pas titrable, le libellé reste le nôtre.
    assert gab.libelle("identification") == "Identification"


def test_un_gabarit_relu_du_disque_est_revalide():
    """Le profil est un fichier éditable, sur un montage synchronisé.

    Relire sans filtrer laisserait entrer par le disque ce que la promesse
    interdit d'y mettre par le réseau.
    """
    truque = Gabarit.depuis_dict({
        "blocs": ["postes", "bloc_inexistant"],
        "intitules": {"postes": "Travaux chez Mme Petitjean",
                      "montants": "montants"},
    })
    # Le bloc inventé tombe ; l'intitulé recopié à la main tombe ; et
    # « montants » aussi, faute d'être dans les blocs retenus.
    assert truque.blocs == ("postes",)
    assert truque.intitules == {}


# --------------------------------------------------------------------------
# 2. Un dépôt refusable ne sort jamais de la machine
# --------------------------------------------------------------------------

def test_un_fichier_vide_est_refuse_et_dit_quoi_faire():
    with pytest.raises(DepotInvalide) as erreur:
        valider_depot(b"")
    assert "enregistr" in str(erreur.value)


def test_un_fichier_qui_n_est_ni_pdf_ni_image_est_refuse():
    with pytest.raises(DepotInvalide) as erreur:
        valider_depot(b"PK\x03\x04 un .docx")
    message = str(erreur.value)
    assert "JPEG" in message and "PNG" in message


def test_le_type_est_lu_au_contenu_pas_a_l_extension():
    assert valider_depot(PDF) == "application/pdf"
    assert valider_depot(JPEG) == "image/jpeg"
    assert valider_depot(PNG) == "image/png"


def test_un_fichier_trop_gros_est_refuse_avec_son_poids():
    gros = b"%PDF-1.7\n" + b"0" * mod.TAILLE_MAX_OCTETS
    with pytest.raises(DepotInvalide) as erreur:
        valider_depot(gros)
    assert "Mo" in str(erreur.value)


def test_aucun_appel_n_est_emis_pour_un_depot_refuse():
    appels: list[tuple] = []

    def espion(octets, type_mime, cle):
        appels.append((type_mime, cle))
        return reponse_ocr(DEVIS_MARKDOWN)

    with pytest.raises(DepotInvalide):
        extraire(b"PK\x03\x04", transport=espion, cle="cle-de-test")
    assert appels == []


# --------------------------------------------------------------------------
# 3. Un échec de lecture ne dégrade pas l'export
# --------------------------------------------------------------------------

def test_sans_cle_le_depot_bascule_sur_le_gabarit_par_defaut(monkeypatch):
    monkeypatch.delenv(mod.VARIABLE_CLE, raising=False)
    resultat = extraire(PDF)
    assert isinstance(resultat, Indisponible)
    assert resultat.motif == "cle_absente"
    assert "défaut" in resultat.explication


def test_un_transport_qui_explose_devient_un_motif_pas_une_exception():
    def en_panne(octets, type_mime, cle):
        raise OSError("connexion refusée vers https://api.example/v1/ocr?k=SECRET")

    resultat = extraire(PDF, transport=en_panne, cle="cle-de-test")
    assert isinstance(resultat, Indisponible)
    assert resultat.motif == "injoignable"
    assert "SECRET" not in resultat.explication


def test_une_photo_illisible_est_nommee_comme_telle():
    """Scénario 2 de la story : dire ce qui n'a pas été lu, et quoi faire."""
    resultat = extraire(JPEG, transport=lambda o, t, c: reponse_ocr("Dev1s ~~ 3"),
                        cle="cle-de-test")
    assert isinstance(resultat, Indisponible)
    assert resultat.motif == "illisible"
    assert "floue" in resultat.explication


def test_un_score_de_confiance_bas_vaut_illisible():
    resultat = extraire(
        JPEG, transport=lambda o, t, c: reponse_ocr(DEVIS_MARKDOWN, confiance=0.2),
        cle="cle-de-test")
    assert isinstance(resultat, Indisponible)
    assert resultat.motif == "illisible"


def test_un_score_absent_n_est_pas_un_mauvais_score():
    """Les modèles anciens n'en renvoient pas : on ne conclut pas d'un chiffre
    inventé, on s'en remet à la longueur du texte."""
    resultat = extraire(
        PDF, transport=lambda o, t, c: reponse_ocr(DEVIS_MARKDOWN, confiance=None),
        cle="cle-de-test")
    assert isinstance(resultat, Gabarit)


def test_une_reponse_sans_pages_devient_un_motif():
    resultat = extraire(PDF, transport=lambda o, t, c: {"erreur": "oups"},
                        cle="cle-de-test")
    assert isinstance(resultat, Indisponible)
    assert resultat.motif == "inattendu"


def test_le_vrai_transport_traduit_les_codes_http_sans_citer_l_url(monkeypatch):
    import urllib.error

    def erreur_http(code):
        def leve(requete, timeout=None):
            raise urllib.error.HTTPError(
                mod.POINT_ENTREE, code, "non", {}, None)
        return leve

    for code, attendu in ((401, "quota"), (429, "quota"), (413, "trop_gros"),
                          (422, "trop_gros"), (500, "injoignable")):
        monkeypatch.setattr(mod.urllib.request, "urlopen", erreur_http(code))
        resultat = extraire(PDF, cle="cle-de-test")
        assert isinstance(resultat, Indisponible)
        assert resultat.motif == attendu, code
        assert "Bearer" not in resultat.explication


def test_chaque_motif_a_son_explication():
    """Un motif sans phrase lisible ferait planter la route sur un KeyError."""
    for motif in mod._EXPLICATIONS:
        assert mod._indisponible(motif).explication


# --------------------------------------------------------------------------
# 4. Un document qui n'est pas un devis est refusé, et on dit pourquoi
# --------------------------------------------------------------------------

def test_une_lettre_d_entreprise_n_est_pas_un_devis():
    lettre = ("# Entreprise Duverger\n\n**Destinataire**\nMme Petitjean\n\n"
              "Madame, nous vous confirmons notre passage. " + "Texte. " * 40)
    with pytest.raises(PasUnDevis) as erreur:
        extraire(PDF, transport=lambda o, t, c: reponse_ocr(lettre),
                 cle="cle-de-test")
    message = erreur.value.explication
    assert "détail des postes" in message and "montants" in message


def test_le_refus_nomme_ce_qui_a_ete_reconnu():
    """Un refus opaque n'aide personne : on dit ce qui manquait."""
    erreur = PasUnDevis(("prestataire", "client"))
    assert "prestataire" in erreur.explication
    assert "client et chantier" in erreur.explication


def test_un_document_muet_est_refuse_sans_enumerer_du_vide():
    erreur = PasUnDevis(())
    assert "ne ressemble pas à un devis" in erreur.explication


def test_des_postes_sans_montants_restent_un_devis():
    """Un bloc décisif suffit, pourvu qu'il y en ait trois en tout."""
    texte = ("## Devis n° 12\n## Prestataire\n## Désignation\n"
             "Peinture des murs du salon, deux couches. " + "Détail. " * 30)
    resultat = extraire(PDF, transport=lambda o, t, c: reponse_ocr(texte),
                        cle="cle-de-test")
    assert isinstance(resultat, Gabarit)


# --------------------------------------------------------------------------
# 5. Le gabarit réordonne et renomme, il ne retire jamais
# --------------------------------------------------------------------------

def test_un_bloc_absent_du_devis_depose_reste_emis():
    gab = Gabarit(blocs=("postes", "montants"), intitules={})
    assert set(gab.ordre()) == set(BLOCS)
    # Ceux du gabarit d'abord, les autres derrière, dans l'ordre par défaut.
    assert gab.ordre()[:2] == ("postes", "montants")


def test_l_ordre_du_devis_depose_est_repris():
    gab = analyser("## Totaux\n## Prestations\n## Prestataire\n")
    assert gab.blocs == ("montants", "postes", "prestataire")


def test_l_intitule_le_plus_long_l_emporte():
    """Sinon « détail des postes » se verrait raccourci en « détail »."""
    gab = analyser("## Détail des postes\n## Prestataire\n## Montants\n")
    assert gab.intitules["postes"] == "détail des postes"


def test_les_titres_en_majuscules_sont_reconnus():
    """Beaucoup de devis d'artisan n'ont pas de titres Markdown."""
    gab = analyser("PRESTATAIRE\nSARL X\nDETAIL DES POSTES\nligne\nTOTAUX\n")
    assert set(gab.blocs) >= {"prestataire", "postes", "montants"}


@pytest.fixture
def profil_complet():
    """Un profil, un client et un devis complets — modèle de `test_export.py`."""
    import datetime as dt

    from app.devis import Devis, ajouter_ligne
    from app.profil import Client, Decennale, ProfilArtisan

    profil = ProfilArtisan(
        raison_sociale="Karim Rénovation SARL", siret="81234567800017",
        adresse="12 rue des Lilas, 69003 Lyon", rcs="RCS Lyon 812 345 678",
        ape_naf="4334Z", tva_intracommunautaire="FR40812345678",
        decennale=Decennale(assureur="AXA", numero_contrat="DEC-99812",
                            couverture_geographique="France métropolitaine",
                            valide_jusqu_au=dt.date(2027, 6, 30)),
    )
    client = Client(nom="Mme Dupont",
                    adresse_chantier="4 impasse du Parc, 69100 Villeurbanne")
    devis = Devis()
    ajouter_ligne(devis, "PEINT_FINITION", "surface", 22.0)
    return devis, profil, client


def test_le_docx_reprend_l_ordre_et_les_intitules(profil_complet):
    from app.export import generer_docx
    devis, profil, client = profil_complet
    gab = Gabarit(blocs=("postes", "montants", "prestataire"),
                  intitules={"postes": "prestations", "montants": "totaux"})
    doc = Document(io.BytesIO(generer_docx(
        devis, [], profil, client, "3 semaines", gabarit=gab)))
    titres = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert titres[1:4] == ["Prestations", "Totaux", "Prestataire"]
    # Aucune section perdue : le devis reste opposable.
    assert "Accord du client" in titres
    assert "Client et chantier" in titres


def test_sans_gabarit_le_docx_ne_change_pas(profil_complet):
    from app.export import generer_docx
    devis, profil, client = profil_complet
    doc = Document(io.BytesIO(generer_docx(devis, [], profil, client, "3 semaines")))
    titres = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert titres[:6] == ["Devis", "Identification", "Prestataire",
                          "Client et chantier", "Détail des postes", "Montants"]


# --------------------------------------------------------------------------
# 6. Le fichier ne passe jamais par le modèle
# --------------------------------------------------------------------------

def test_le_module_n_importe_pas_le_modele():
    """Vérifié sur l'AST, pas sur une chaîne : `app.llm` ne doit pas entrer ici.

    C'est le critère nº 1 de la story, et le seul qu'un test puisse tenir
    structurellement plutôt que par intention.
    """
    arbre = ast.parse((RACINE / "app" / "gabarit.py").read_text(encoding="utf-8"))
    importes = set()
    for noeud in ast.walk(arbre):
        if isinstance(noeud, ast.Import):
            importes.update(a.name for a in noeud.names)
        elif isinstance(noeud, ast.ImportFrom) and noeud.module:
            importes.add(noeud.module)
    assert not any(nom.startswith(("app.llm", "app.pipeline")) for nom in importes)


def test_la_cle_d_ocr_n_est_pas_ecrite_au_depot():
    """Aucun littéral affecté à un nom de clé dans `app/` (même garde qu'US-01)."""
    suspects = ("cle", "clef", "token", "api_key", "secret", "password",
                "mot_de_passe")
    for fichier in (RACINE / "app").rglob("*.py"):
        arbre = ast.parse(fichier.read_text(encoding="utf-8"))
        for noeud in ast.walk(arbre):
            if not isinstance(noeud, ast.Assign):
                continue
            for cible in noeud.targets:
                nom = getattr(cible, "id", "").lower()
                # `VARIABLE_CLE` nomme une variable d'environnement, pas une
                # clé : c'est la seule exception, et elle est nommée.
                if nom.startswith("variable_"):
                    continue
                if any(s in nom for s in suspects) and isinstance(
                        noeud.value, ast.Constant) and isinstance(
                        noeud.value.value, str):
                    raise AssertionError(f"{fichier}: {nom}")


def test_le_front_ignore_le_nom_du_service_d_ocr():
    """Le navigateur n'a pas à savoir qui lit le document, ni avec quelle clé."""
    html = (RACINE / "app" / "static" / "index.html").read_text(encoding="utf-8")
    for interdit in ("mistral", "api.mistral.ai", "Bearer", mod.VARIABLE_CLE):
        assert interdit.lower() not in html.lower()


# --------------------------------------------------------------------------
# 7. La route : dépôt, repli, révocation
# --------------------------------------------------------------------------

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "REGISTRE", Registre(racine=tmp_path))
    monkeypatch.setattr(pipeline, "appeler_modele",
                        lambda h, m, timeout=60: {"intent": "reponse",
                                                  "mention": None,
                                                  "payload": {"texte": "ok"}})
    c = TestClient(main.app)
    authentifier(c)
    return c


@pytest.fixture
def ocr_qui_repond(monkeypatch):
    monkeypatch.setenv(mod.VARIABLE_CLE, "cle-de-test")
    monkeypatch.setattr(mod, "_transport_http",
                        lambda octets, type_mime, cle: reponse_ocr(DEVIS_MARKDOWN))


def deposer(client, contenu=PDF, nom="devis.pdf", type_mime="application/pdf"):
    return client.post("/api/profil/gabarit",
                       files={"fichier": (nom, contenu, type_mime)})


def test_la_route_retient_la_mise_en_page(client, ocr_qui_repond):
    r = deposer(client)
    assert r.status_code == 200
    corps = r.json()
    assert corps["origine"] == "devis_depose"
    assert corps["motif"] is None
    assert corps["gabarit"]["par_defaut"] is False
    assert corps["gabarit"]["blocs"][:3] == ["identification", "prestataire",
                                             "client"]


def test_la_route_dit_que_les_sections_manquantes_restent(client, ocr_qui_repond):
    corps = deposer(client).json()
    assert "restent ajoutées" in corps["explication"]


def test_un_depot_refuse_repond_422_et_designe_le_champ(client, ocr_qui_repond):
    r = deposer(client, contenu=b"PK\x03\x04", nom="devis.docx")
    assert r.status_code == 422
    assert r.json()["detail"]["champ"] == "fichier"


def test_un_document_qui_n_est_pas_un_devis_repond_200(client, monkeypatch):
    """Le produit n'est pas en panne : c'est le document qui ne convient pas."""
    monkeypatch.setenv(mod.VARIABLE_CLE, "cle-de-test")
    monkeypatch.setattr(mod, "_transport_http",
                        lambda o, t, c: reponse_ocr("# Facture\n" + "Texte. " * 60))
    r = deposer(client)
    assert r.status_code == 200
    assert r.json()["motif"] == "pas_un_devis"
    assert r.json()["gabarit"]["par_defaut"] is True


def test_une_panne_d_ocr_repond_200_et_garde_le_defaut(client, monkeypatch):
    monkeypatch.delenv(mod.VARIABLE_CLE, raising=False)
    r = deposer(client)
    assert r.status_code == 200
    assert r.json()["motif"] == "cle_absente"
    assert r.json()["gabarit"]["par_defaut"] is True


def test_le_gabarit_survit_a_un_redemarrage(client, ocr_qui_repond):
    deposer(client)
    projet_id = client.cookies.get("projet_id")
    main.REGISTRE._cache.clear()          # ce que fait un redémarrage du serveur
    session = main.REGISTRE.obtenir(projet_id)
    assert session.profil.gabarit.blocs
    assert session.profil.gabarit.intitules["postes"] == "prestations"


def test_la_mise_en_page_est_revocable(client, ocr_qui_repond):
    deposer(client)
    r = client.delete("/api/profil/gabarit")
    assert r.status_code == 200
    assert r.json()["gabarit"]["par_defaut"] is True
    assert client.get("/api/profil").json()["gabarit"]["par_defaut"] is True


def test_revoquer_sans_rien_a_revoquer_reussit(client):
    """Demander l'oubli et s'entendre répondre 404 serait absurde."""
    assert client.delete("/api/profil/gabarit").status_code == 200


def test_le_journal_ne_porte_ni_nom_de_fichier_ni_contenu(client, ocr_qui_repond):
    from app.observabilite import JOURNAL
    deposer(client, nom="devis-Petitjean-2024.pdf")
    trace = str(JOURNAL.evenements())
    assert "Petitjean" not in trace
    assert "devis-Petitjean-2024.pdf" not in trace


def test_le_depot_exige_une_session():
    anonyme = TestClient(main.app)
    r = anonyme.post("/api/profil/gabarit",
                     files={"fichier": ("d.pdf", PDF, "application/pdf")})
    assert r.status_code == 401


# --------------------------------------------------------------------------
# Ce que ces tests ne prouvent pas
# --------------------------------------------------------------------------
#
# Aucun appel réel n'a été passé au service d'OCR : aucune clé n'est
# configurée sur cette machine. La ligne de requête et la forme de la réponse
# (`pages[].markdown`, `confidence_scores.average_page_confidence_score`)
# viennent de la documentation publiée, **pas d'une réponse observée**.
#
# Ce qui est prouvé ici, c'est la **tolérance** du code : quelle que soit la
# forme réelle de la réponse, une clé manquante, un champ absent, un score
# inattendu ou une panne se résolvent en un motif et un repli, jamais en une
# trace d'exception ni en un gabarit inventé. Et surtout : quoi que renvoie le
# service, rien de ce qu'il renvoie n'entre dans le gabarit.
